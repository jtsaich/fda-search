
import asyncio
import io
import openpyxl
from docx import Document
from services.document_service import DocumentService

async def test_extraction():
    service = DocumentService()

    # Test 1: CSV Extraction
    print("Testing CSV extraction...")
    csv_content = "header1,header2\nvalue1,value2\nvalue3,value4"
    csv_bytes = csv_content.encode('utf-8')

    extracted_csv = await service.extract_text_from_bytes(csv_bytes, "test.csv")
    print(f"CSV Extracted: {extracted_csv}")
    assert "header1, header2" in extracted_csv
    assert "value1, value2" in extracted_csv

    # Test 2: XLSX Extraction
    print("\nTesting XLSX extraction...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["col1", "col2"])
    ws.append(["data1", "data2"])

    xlsx_io = io.BytesIO()
    wb.save(xlsx_io)
    xlsx_bytes = xlsx_io.getvalue()

    extracted_xlsx = await service.extract_text_from_bytes(xlsx_bytes, "test.xlsx")
    print(f"XLSX Extracted: {extracted_xlsx}")
    assert "col1, col2" in extracted_xlsx
    assert "data1, data2" in extracted_xlsx

    # Test 3: DOCX Extraction
    print("\nTesting DOCX extraction...")
    doc = Document()
    doc.add_paragraph("Hello World DOCX")

    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    extracted_docx = await service.extract_text_from_bytes(docx_bytes, "test.docx")
    print(f"DOCX Extracted: {extracted_docx}")
    assert "Hello World DOCX" in extracted_docx

    print("\nAll tests passed!")

if __name__ == "__main__":
    asyncio.run(test_extraction())
